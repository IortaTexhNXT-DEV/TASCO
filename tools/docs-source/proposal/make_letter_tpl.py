"""Cover letter on the iorta TechNXT Word template: keeps the template's page setup, header and
footer (logo and office addresses), drops the report frame (cover, contents, control tables)."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
src, tpl, out = sys.argv[1], sys.argv[2], sys.argv[3]
NAVY = RGBColor(0x1F, 0x2A, 0x5C)
d = Document(tpl)
body = d.element.body
for el in list(body):
    if not el.tag.endswith('}sectPr'):
        body.remove(el)

def para(text='', bold=False, size=10, after=5, color=None):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.bold = bold or i % 2 == 1
        if color is not None:
            r.font.color.rgb = color
    return p

lines = open(src, encoding='utf-8').read().split('\n')
lines = [l for l in lines if not l.startswith('# ')]
block = []
def flush():
    if not block:
        return
    text = ' '.join(block) if not all(b and not b.endswith('.') for b in block[:1]) else None
    for b in block if text is None or len(block) > 1 and not block[0].endswith(('.', ',')) else [' '.join(block)]:
        b = b.replace('&nbsp;', '').strip()
        if b.startswith('**Subject:'):
            para(b, color=NAVY, after=10)
        else:
            para(b, after=0 if len(block) > 1 else 6)
    if len(block) > 1:
        para('', after=4)
    block.clear()
for l in lines:
    if l.strip() == '':
        flush()
    else:
        block.append(l.rstrip())
flush()
d.save(out)
print('ok', out)
