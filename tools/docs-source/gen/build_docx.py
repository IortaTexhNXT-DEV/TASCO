#!/usr/bin/env python3
"""Markdown -> Word generator in the iorta TechNXT corporate template.

Usage:
    build_docx.py <input.md> <output.docx> [--img-dir DIR] [--template T.docx]
                  [--toc-mode inject|roundtrip|none] [--preview DIR] [--quiet]

See README.txt in this folder for the full description.
"""
import argparse
import copy
import datetime as dt
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import yaml
from lxml import etree
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SCRATCH = os.path.dirname(HERE)
DEFAULT_TEMPLATE = os.path.join(SCRATCH, 'tpl', 'template.docx')
MERMAID_JS = os.environ.get('MERMAID_JS') or os.path.join(SCRATCH, 'vendor', 'mermaid.min.js')
RENDER_JS = os.path.join(HERE, 'render_mermaid.js')
LUA_FILTER = os.path.join(HERE, 'filter.lua')
LO_TOOLS = os.path.join(HERE, 'lo_tools.py')
CACHE_DIR = os.path.join(HERE, '.cache', 'mermaid')
UNO_PYTHON = '/usr/bin/python3'

# ---------------------------------------------------------------- constants
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {
    'w': W_NS,
    'r': R_NS,
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape',
    'v': 'urn:schemas-microsoft-com:vml',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'ct': 'http://schemas.openxmlformats.org/package/2006/content-types',
    'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
    'dc': 'http://purl.org/dc/elements/1.1/',
    'dcterms': 'http://purl.org/dc/terms/',
    'ep': 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties',
}
REL_IMAGE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'
REL_HYPERLINK = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'
REL_FOOTER = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer'
REL_CUSTOMXML = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml'
REL_CUSTOMPROPS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties'

HEAD_BLUE = '0F4761'     # template heading colour (Heading 1 text)
PRIMARY = '1D74BA'       # template cover/title blue
TEXT = '101820'          # template body text colour
GRID = '808080'          # template table border grey
BAND = 'F4F7FB'          # zebra band for body tables
CALLOUT_FILL = 'E8F1FA'

PAGE_W, PAGE_H = 11907, 16839
MARGIN = 1440
PORTRAIT_TW = PAGE_W - 2 * MARGIN        # 9027 twips text width
LANDSCAPE_TW = PAGE_H - 2 * MARGIN       # 13959 twips
EMU_PER_CM = 360000

REQUIREMENT_DOC_IDS = {'TGP-BUS-02', 'TGP-BUS-03', 'TGP-BUS-04'}

MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
          'September', 'October', 'November', 'December']


def w(tag):
    p, t = tag.split(':')
    return '{%s}%s' % (NS[p], t)


def wattr(el, name, default=None):
    return el.get(w(name), default) if el is not None else default


def E(tag, attrs=None, *children):
    el = etree.Element(w(tag))
    for k, v in (attrs or {}).items():
        el.set(w(k) if ':' in k else k, str(v))
    for c in children:
        if c is not None:
            el.append(c)
    return el


def frag(xml):
    """Parse a WordprocessingML fragment written with the w:/r:/... prefixes."""
    decl = ' '.join('xmlns:%s="%s"' % (k, v) for k, v in NS.items() if k in ('w', 'r', 'wp', 'a', 'pic'))
    return etree.fromstring('<w:root %s>%s</w:root>' % (decl, xml))[0]


def text_of(el):
    return ''.join(t.text or '' for t in el.iter(w('w:t')))


def log(msg, quiet=False):
    if not quiet:
        print(msg, file=sys.stderr)


WARNINGS = []


def warn(msg):
    WARNINGS.append(msg)
    print('WARNING: ' + msg, file=sys.stderr)


# ------------------------------------------------------------ front matter
FM_RE = re.compile(r'^\ufeff?---[ \t]*\r?\n(.*?)\r?\n(?:---|\.\.\.)[ \t]*\r?\n', re.S)


def split_front_matter(text):
    m = FM_RE.match(text)
    if not m:
        return None, text
    data = yaml.safe_load(m.group(1)) or {}
    if not isinstance(data, dict):
        return None, text
    return data, text[m.end():]


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value or '').strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d %B %Y', '%d %b %Y', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def as_list(item):
    if isinstance(item, (list, tuple)):
        return [str(x).strip() for x in item]
    if isinstance(item, dict):
        return [str(x).strip() for x in item.values()]
    return [s.strip() for s in str(item).split('|')]


def normalise_meta(fm):
    m = {}
    m['id'] = str(fm.get('id', '')).strip()
    m['title'] = str(fm.get('title', '')).strip()
    m['subtitle'] = str(fm.get('subtitle', '')).strip()
    v = fm.get('version', '1.0')
    m['version'] = str(v).strip() if v is not None else '1.0'
    d = parse_date(fm.get('date'))
    m['date'] = d
    m['date_short'] = d.strftime('%d/%m/%Y') if d else str(fm.get('date', '')).strip()
    m['date_long'] = ('%02d %s %d' % (d.day, MONTHS[d.month - 1], d.year)) if d else m['date_short']
    m['prepared_by'] = str(fm.get('prepared_by', 'iorta TechNXT')).strip()
    m['reviewed_by'] = str(fm.get('reviewed_by', '')).strip()
    m['approved_by'] = str(fm.get('approved_by', '')).strip()
    m['change_history'] = str(fm.get('change_history', 'Initial issue')).strip()
    acr = []
    for item in fm.get('acronyms') or []:
        parts = as_list(item)
        if not parts or not parts[0]:
            continue
        acr.append((parts[0], ', '.join(p for p in parts[1:] if p)))
    seen = {}
    for a, mean in acr:
        seen.setdefault(a, mean)
    m['acronyms'] = sorted(seen.items(), key=lambda x: (x[0].lower(), x[0]))
    so = []
    for item in fm.get('signoff') or []:
        parts = as_list(item)
        if len(parts) >= 3:
            so.append((', '.join(parts[:-2]), parts[-2], parts[-1]))
        elif len(parts) == 2:
            so.append((parts[0], parts[1], 'Open'))
        elif parts and parts[0]:
            so.append((parts[0], '', 'Open'))
    m['signoff'] = so
    return m


def split_org(s):
    """'iorta TechNXT, Solution Architecture' -> ('iorta TechNXT', 'Solution Architecture')."""
    if ',' in s:
        a, b = s.split(',', 1)
        return a.strip(), b.strip()
    return s.strip(), ''


# ------------------------------------------------------- markdown handling
FENCE_RE = re.compile(r'^(\s{0,3})(`{3,}|~{3,})\s*([^`\s]*)?.*$')
HEADING_RE = re.compile(r'^(#{1,6})\s+(.*?)\s*#*\s*$')
NUM_PREFIX_RE = re.compile(r'^(?:\d+(?:\.\d+)*\.?|[0-9]+\))\s+(?=\S)')
TABLE_LINE_RE = re.compile(r'^\s*\|')
IMG_LINE_RE = re.compile(r'^\s*!\[(?P<alt>[^\]]*)\]\((?P<src>[^)\s]+)(?:\s+"(?P<title>[^"]*)")?\)\s*(?:\{(?P<attrs>[^}]*)\})?\s*$')
TABLE_CAP_RE = re.compile(r'^\s*(?:Table|:)\s*:?\s+(?P<cap>.+?)\s*$')


def fence_mask(lines):
    """Return list of booleans: True when the line is inside (or delimits) a fenced block."""
    mask = []
    fence = None
    for ln in lines:
        m = FENCE_RE.match(ln)
        if fence is None:
            if m:
                fence = m.group(2)[0] * len(m.group(2))
                mask.append(True)
            else:
                mask.append(False)
        else:
            mask.append(True)
            if m and m.group(2).startswith(fence) and not (m.group(3) or '').strip():
                fence = None
    return mask


def normalise_markdown(body, meta):
    """Accept older drafts: a single H1 title, a metadata table and typed heading numbers."""
    lines = body.split('\n')
    mask = fence_mask(lines)
    h1 = [i for i, ln in enumerate(lines) if not mask[i] and re.match(r'^#\s+\S', ln)]
    first = next((i for i, ln in enumerate(lines) if ln.strip()), None)
    def _n(x):
        return re.sub(r'[^a-z0-9]+', '', x.lower())
    if (first is not None and len(h1) == 1 and h1[0] == first and meta.get('title')
            and _n(meta['title']) in _n(lines[first])):
        warn('body starts with an H1 title; it was dropped and headings were promoted one level')
        del lines[first]
        del mask[first]
        for i, ln in enumerate(lines):
            if not mask[i] and re.match(r'^#{2,6}\s', ln):
                lines[i] = ln[1:]
        # metadata table straight after the title (| Document | Version | ...)
        j = first
        while j < len(lines) and not lines[j].strip():
            j += 1
        if j < len(lines) and TABLE_LINE_RE.match(lines[j]) and re.search(r'\bDocument\b.*\bVersion\b', lines[j]):
            k = j
            while k < len(lines) and TABLE_LINE_RE.match(lines[k]):
                k += 1
            del lines[j:k]
            del mask[j:k]
    out = []
    for i, ln in enumerate(lines):
        if not mask[i]:
            m = HEADING_RE.match(ln)
            if m:
                txt = NUM_PREFIX_RE.sub('', m.group(2))
                ln = '%s %s' % (m.group(1), txt)
        out.append(ln)
    return '\n'.join(out)


def png_size_cm(path, scale=1.0):
    with Image.open(path) as im:
        wpx, hpx = im.size
        dpi = im.info.get('dpi', (96, 96))[0] or 96
    if scale != 1.0:
        dpi = 96 * scale
    return wpx / dpi * 2.54, hpx / dpi * 2.54


def fit_image(w_cm, h_cm, allow_landscape=True):
    """Return (width_cm, height_cm, landscape?) choosing the orientation that shows the image largest."""
    fp = min(1.0, 16.5 / w_cm, 18.0 / h_cm)
    fl = min(1.0, 24.0 / w_cm, 14.0 / h_cm)
    landscape = allow_landscape and fl > fp * 1.15
    f = fl if landscape else fp
    return w_cm * f, h_cm * f, landscape


def single_sentence(text):
    t = text.strip()
    if not t or len(t) > 260:
        return False
    core = t.rstrip('.:')
    return not re.search(r'[.!?]\s+[A-Z(]', core)


def render_jobs(jobs, work, quiet=False):
    todo = [j for j in jobs if not os.path.isfile(j['out'])]
    if not todo:
        return
    os.makedirs(CACHE_DIR, exist_ok=True)
    jf = os.path.join(work, 'mermaid_jobs.json')
    with open(jf, 'w', encoding='utf-8') as f:
        json.dump([{'code': j['code'], 'out': j['out']} for j in todo], f)
    res = subprocess.run(['node', RENDER_JS, jf, MERMAID_JS], capture_output=True, text=True, timeout=600)
    try:
        summary = json.loads(res.stdout.strip().splitlines()[-1])
    except Exception:
        summary = {'errors': [{'out': '*', 'message': res.stderr.strip()[-400:]}]}
    for e in summary.get('errors', []):
        warn('mermaid render failed (%s): %s' % (os.path.basename(e['out']), e['message']))
    log('mermaid: rendered %s, failed %s' % (summary.get('rendered'), summary.get('failed')), quiet)


def preprocess(body, md_dir, img_dir, work, quiet=False):
    """Mermaid -> PNG figures, images -> numbered figures, 'Table:' -> numbered table captions."""
    lines = body.split('\n')
    out = []
    jobs = []
    fig_n = 0
    tab_n = 0
    last_heading = ''
    i = 0
    fence = None

    def prev_paragraph():
        """(start, end) indices in out of the paragraph directly before the current point."""
        end = len(out)
        while end > 0 and not out[end - 1].strip():
            end -= 1
        if end == 0:
            return None
        start = end
        while start > 0 and out[start - 1].strip():
            start -= 1
        para = out[start:end]
        first = para[0].lstrip()
        if re.match(r'^(#|\||>|[-*+]\s|\d+[.)]\s|:::|!\[|```|~~~)', first):
            return None
        return start, end

    def emit_figure(path, caption, w_cm, h_cm, landscape):
        nonlocal fig_n
        fig_n += 1
        style = 'FigureWide' if landscape else 'Figure'
        alt = caption.replace('[', '(').replace(']', ')')
        out.extend(['', '::: {custom-style="%s"}' % style,
                    '![%s](%s){width=%.2fcm height=%.2fcm}' % (alt, path, w_cm, h_cm), ':::', ''])
        cap = 'Figure %d \u2013 %s' % (fig_n, caption) if caption else 'Figure %d' % fig_n
        out.extend(['::: {custom-style="Caption"}', cap, ':::', ''])

    while i < len(lines):
        ln = lines[i]
        m = FENCE_RE.match(ln)
        if fence is None and m:
            info = (m.group(3) or '').strip().lower()
            fchar = m.group(2)
            j = i + 1
            while j < len(lines):
                mm = FENCE_RE.match(lines[j])
                if mm and mm.group(2).startswith(fchar) and not (mm.group(3) or '').strip():
                    break
                j += 1
            block = lines[i + 1:j]
            if info == 'mermaid':
                caption = None
                code_lines = []
                for cl in block:
                    cm = re.match(r'^\s*%%\s*caption\s*:\s*(.+?)\s*$', cl, re.I)
                    if cm:
                        caption = cm.group(1)
                    else:
                        code_lines.append(cl)
                if caption is None:
                    pp = prev_paragraph()
                    if pp:
                        ptxt = ' '.join(x.strip() for x in out[pp[0]:pp[1]])
                        if single_sentence(ptxt):
                            caption = ptxt
                            del out[pp[0]:pp[1]]
                if caption is None:
                    caption = last_heading
                caption = re.sub(r'\s*[:.]\s*$', '', caption or '').strip()
                code = '\n'.join(code_lines).strip() + '\n'
                key = hashlib.sha1((code + open(RENDER_JS, encoding='utf-8').read()).encode()).hexdigest()[:16]
                png = os.path.join(CACHE_DIR, key + '.png')
                jobs.append({'code': code, 'out': png, 'placeholder': len(out), 'caption': caption})
                out.append('\x00MERMAID%d\x00' % (len(jobs) - 1))
                i = j + 1
                continue
            # ordinary code fence: copy through
            if out and out[-1].strip():
                out.append('')
            out.extend(lines[i:j + 1])
            i = j + 1
            continue
        hm = HEADING_RE.match(ln)
        if hm:
            last_heading = re.sub(r'[*_`]', '', hm.group(2)).strip()
            if out and out[-1].strip():
                out.append('')
            out.append(ln)
            i += 1
            continue
        im = IMG_LINE_RE.match(ln)
        if im:
            src = im.group('src')
            caption = (im.group('title') or im.group('alt') or '').strip()
            if re.match(r'^https?://', src):
                warn('remote image not embedded: %s' % src)
                out.append('*[Image: %s]*' % (caption or src))
                i += 1
                continue
            cands = [os.path.join(md_dir, src)]
            if img_dir:
                cands += [os.path.join(img_dir, src), os.path.join(img_dir, os.path.basename(src))]
            path = next((c for c in cands if os.path.isfile(c)), None)
            if not path:
                warn('image not found: %s' % src)
                out.append('*[Missing image: %s]*' % (caption or src))
                i += 1
                continue
            path = os.path.abspath(path)
            if path.lower().endswith('.svg'):
                warn('SVG images are not supported by every Word version; convert to PNG: %s' % src)
            try:
                wc, hc = png_size_cm(path)
            except Exception:
                wc, hc = 16.5, 10.0
            wm = re.search(r'width\s*=\s*"?([\d.]+)cm', im.group('attrs') or '')
            if wm:
                # An explicit width keeps the figure on a portrait page at that size (capped to the text block).
                f = min(float(wm.group(1)), 16.5) / wc
                f = min(f, 18.0 / hc)
                wc, hc, ls = wc * f, hc * f, False
            else:
                wc, hc, ls = fit_image(wc, hc)
            emit_figure(path, caption, wc, hc, ls)
            i += 1
            continue
        tm = re.match(r'^\s*Table:\s*(?P<cap>.+?)\s*$', ln)
        if tm:
            cap_text = tm.group('cap')
            # caption after a table?
            k = len(out)
            while k > 0 and not out[k - 1].strip():
                k -= 1
            if k > 0 and TABLE_LINE_RE.match(out[k - 1]):
                s = k - 1
                while s > 0 and TABLE_LINE_RE.match(out[s - 1]):
                    s -= 1
                tab_n += 1
                out[s:s] = ['::: {custom-style="TableCaption"}', 'Table %d \u2013 %s' % (tab_n, cap_text), ':::', '']
            else:
                tab_n += 1
                out.extend(['', '::: {custom-style="TableCaption"}', 'Table %d \u2013 %s' % (tab_n, cap_text), ':::', ''])
            i += 1
            continue
        if TABLE_LINE_RE.match(ln) and out and out[-1].strip() and not TABLE_LINE_RE.match(out[-1]) \
                and not out[-1].startswith(':::'):
            out.append('')
        if ln.lstrip().startswith('>') and out and out[-1].strip() and not out[-1].lstrip().startswith('>'):
            out.append('')
        out.append(ln)
        i += 1

    render_jobs(jobs, work, quiet)
    # Very wide left-to-right strips become unreadable when shrunk to the page; try top-down too.
    alts = []
    for j in jobs:
        if not os.path.isfile(j['out']):
            continue
        wc, hc = png_size_cm(j['out'], scale=2.0)
        fw, fh, _ = fit_image(wc, hc)
        scale = fw / wc
        j['scale'] = scale
        first = j['code'].lstrip().split('\n', 1)[0]
        if scale < 0.65 and wc / hc > 3 and re.match(r'^(flowchart|graph)\s+(LR|RL)\b', first):
            code = re.sub(r'^(\s*(?:flowchart|graph)\s+)(LR|RL)', r'\1TB', j['code'], count=1)
            key = hashlib.sha1((code + open(RENDER_JS, encoding='utf-8').read()).encode()).hexdigest()[:16]
            alt = {'code': code, 'out': os.path.join(CACHE_DIR, key + '.png'), 'parent': j}
            alts.append(alt)
    render_jobs(alts, work, quiet)
    for alt in alts:
        j = alt['parent']
        if os.path.isfile(alt['out']):
            wc, hc = png_size_cm(alt['out'], scale=2.0)
            fw, fh, _ = fit_image(wc, hc)
            if fw / wc > j['scale'] * 1.25:
                warn('diagram "%s" is a long left-to-right strip; rendered top-down for legibility' % j['caption'])
                j['out'] = alt['out']
                j['scale'] = fw / wc
    for j in jobs:
        if j.get('scale', 1) < 0.6:
            warn('diagram "%s" is shrunk to %d%% to fit the page; its text may be small (simplify it)'
                 % (j['caption'], round(j['scale'] * 100)))

    # substitute placeholders (in order, so figure numbers follow the document)
    final = []
    marker = re.compile('^\x00MERMAID(\\d+)\x00$')
    saved_out = out
    out = final
    for ln in saved_out:
        mm = marker.match(ln)
        if not mm:
            final.append(ln)
            continue
        job = jobs[int(mm.group(1))]
        if os.path.isfile(job['out']):
            wc, hc = png_size_cm(job['out'], scale=2.0)
            wc, hc, ls = fit_image(wc, hc)
            emit_figure(job['out'], job['caption'], wc, hc, ls)
        else:
            final.extend(['', '```text'] + job['code'].rstrip('\n').split('\n') + ['```', ''])
    # renumber figures in document order (images and mermaid were numbered in two passes)
    n = 0
    for idx, ln in enumerate(final):
        if re.match('^Figure \\d+( \u2013 |$)', ln) and idx > 0 and final[idx - 1] == '::: {custom-style="Caption"}':
            n += 1
            final[idx] = re.sub(r'^Figure \d+', 'Figure %d' % n, ln)
    return '\n'.join(final), n, tab_n


# --------------------------------------------------------------- styles
def _rfonts(face):
    return '<w:rFonts w:ascii="{0}" w:hAnsi="{0}" w:eastAsia="{0}" w:cs="{0}"/>'.format(face)


SEGOE = _rfonts('Segoe UI')
CONSOLAS = _rfonts('Consolas')
MAJOR = '<w:rFonts w:asciiTheme="majorHAnsi" w:eastAsiaTheme="majorEastAsia" w:hAnsiTheme="majorHAnsi" w:cstheme="majorBidi"/>'

STYLE_DEFS = [
    # paragraph styles used by pandoc
    ('paragraph', 'BodyText', 'Body Text', 'Normal', None,
     '<w:pPr><w:spacing w:before="0" w:after="140" w:line="276" w:lineRule="auto"/><w:jc w:val="both"/></w:pPr>'
     '<w:rPr>%s<w:color w:val="%s"/><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr>' % (SEGOE, TEXT)),
    ('paragraph', 'FirstParagraph', 'First Paragraph', 'BodyText', None, ''),
    ('paragraph', 'Compact', 'Compact', 'BodyText', None,
     '<w:pPr><w:spacing w:before="0" w:after="60"/><w:jc w:val="left"/></w:pPr>'),
    ('paragraph', 'TableText', 'Table Text', 'Normal', None,
     '<w:pPr><w:spacing w:before="20" w:after="20" w:line="245" w:lineRule="auto"/><w:jc w:val="left"/></w:pPr>'
     '<w:rPr>%s<w:color w:val="%s"/><w:sz w:val="19"/><w:szCs w:val="19"/></w:rPr>' % (SEGOE, TEXT)),
    ('paragraph', 'TableHeading', 'Table Heading', 'TableText', None,
     '<w:pPr><w:keepNext/></w:pPr><w:rPr><w:b/><w:bCs/><w:color w:val="FFFFFF"/></w:rPr>'),
    ('paragraph', 'Heading2', 'heading 2', 'Normal', 'BodyText',
     '<w:pPr><w:keepNext/><w:keepLines/><w:numPr><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr>'
     '<w:spacing w:before="280" w:after="100" w:line="259" w:lineRule="auto"/><w:ind w:left="720" w:hanging="720"/>'
     '<w:outlineLvl w:val="1"/></w:pPr>'
     '<w:rPr>%s<w:b/><w:bCs/><w:color w:val="%s"/><w:sz w:val="26"/><w:szCs w:val="26"/></w:rPr>' % (MAJOR, HEAD_BLUE)),
    ('paragraph', 'Heading3', 'heading 3', 'Normal', 'BodyText',
     '<w:pPr><w:keepNext/><w:keepLines/><w:numPr><w:ilvl w:val="2"/><w:numId w:val="1"/></w:numPr>'
     '<w:spacing w:before="220" w:after="80" w:line="259" w:lineRule="auto"/><w:ind w:left="864" w:hanging="864"/>'
     '<w:outlineLvl w:val="2"/></w:pPr>'
     '<w:rPr>%s<w:b/><w:bCs/><w:color w:val="%s"/><w:sz w:val="23"/><w:szCs w:val="23"/></w:rPr>' % (MAJOR, HEAD_BLUE)),
    ('paragraph', 'Heading4', 'heading 4', 'Normal', 'BodyText',
     '<w:pPr><w:keepNext/><w:keepLines/><w:spacing w:before="180" w:after="60" w:line="259" w:lineRule="auto"/>'
     '<w:outlineLvl w:val="3"/></w:pPr>'
     '<w:rPr>%s<w:b/><w:bCs/><w:color w:val="%s"/><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr>' % (SEGOE, HEAD_BLUE)),
    ('paragraph', 'Caption', 'caption', 'Normal', 'BodyText',
     '<w:pPr><w:spacing w:before="60" w:after="240" w:line="240" w:lineRule="auto"/><w:jc w:val="center"/></w:pPr>'
     '<w:rPr>%s<w:i/><w:iCs/><w:color w:val="%s"/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>' % (SEGOE, HEAD_BLUE)),
    ('paragraph', 'TableCaption', 'Table Caption', 'Caption', None,
     '<w:pPr><w:keepNext/><w:spacing w:before="160" w:after="80"/><w:jc w:val="left"/></w:pPr>'),
    ('paragraph', 'Figure', 'Figure', 'Normal', 'Caption',
     '<w:pPr><w:keepNext/><w:spacing w:before="160" w:after="40" w:line="240" w:lineRule="auto"/><w:jc w:val="center"/></w:pPr>'),
    ('paragraph', 'FigureWide', 'Figure Wide', 'Figure', 'Caption', ''),
    ('paragraph', 'CodeBlock', 'Code Block', 'Normal', 'BodyText',
     '<w:pPr><w:pBdr><w:top w:val="single" w:sz="4" w:space="4" w:color="D3DCE6"/>'
     '<w:left w:val="single" w:sz="4" w:space="4" w:color="D3DCE6"/><w:bottom w:val="single" w:sz="4" w:space="4" w:color="D3DCE6"/>'
     '<w:right w:val="single" w:sz="4" w:space="4" w:color="D3DCE6"/></w:pBdr>'
     '<w:shd w:val="clear" w:color="auto" w:fill="F3F6FA"/><w:spacing w:before="80" w:after="200" w:line="240" w:lineRule="auto"/>'
     '<w:ind w:left="113" w:right="113"/><w:jc w:val="left"/></w:pPr>'
     '<w:rPr>%s<w:color w:val="1F2937"/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>' % CONSOLAS),
    ('character', 'CodeChar', 'Code Char', 'DefaultParagraphFont', None,
     '<w:rPr>%s<w:color w:val="1F3864"/><w:sz w:val="19"/><w:szCs w:val="19"/></w:rPr>' % CONSOLAS),
    ('paragraph', 'Callout', 'Callout', 'Normal', 'BodyText',
     '<w:pPr><w:pBdr><w:left w:val="single" w:sz="24" w:space="8" w:color="%s"/></w:pBdr>'
     '<w:shd w:val="clear" w:color="auto" w:fill="%s"/><w:spacing w:before="60" w:after="60" w:line="264" w:lineRule="auto"/>'
     '<w:ind w:left="227" w:right="113"/><w:jc w:val="left"/></w:pPr>'
     '<w:rPr>%s<w:color w:val="%s"/><w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr>' % (PRIMARY, CALLOUT_FILL, SEGOE, TEXT)),
    ('paragraph', 'TOC2', 'toc 2', 'Normal', 'Normal',
     '<w:pPr><w:tabs><w:tab w:val="left" w:pos="1260"/><w:tab w:val="right" w:leader="dot" w:pos="9350"/></w:tabs>'
     '<w:spacing w:after="60" w:line="264" w:lineRule="auto"/><w:ind w:left="360"/></w:pPr>'
     '<w:rPr>%s<w:noProof/><w:color w:val="%s"/><w:sz w:val="19"/><w:szCs w:val="19"/></w:rPr>' % (SEGOE, TEXT)),
    ('paragraph', 'TOC3', 'toc 3', 'Normal', 'Normal',
     '<w:pPr><w:tabs><w:tab w:val="left" w:pos="1700"/><w:tab w:val="right" w:leader="dot" w:pos="9350"/></w:tabs>'
     '<w:spacing w:after="40" w:line="264" w:lineRule="auto"/><w:ind w:left="720"/></w:pPr>'
     '<w:rPr>%s<w:noProof/><w:color w:val="404A55"/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>' % SEGOE),
]


def install_styles(styles_xml, extra_from=None):
    root = etree.fromstring(styles_xml)
    have = {s.get(w('w:styleId')): s for s in root.findall(w('w:style'))}
    for kind, sid, name, based, nxt, body in STYLE_DEFS:
        xml = '<w:style w:type="%s" w:customStyle="1" w:styleId="%s"><w:name w:val="%s"/>' % (kind, sid, name)
        if name.startswith(('heading', 'toc', 'caption')) or sid in ('BodyText',):
            xml = xml.replace(' w:customStyle="1"', '')
        if based:
            xml += '<w:basedOn w:val="%s"/>' % based
        if nxt:
            xml += '<w:next w:val="%s"/>' % nxt
        if sid.startswith('Heading'):
            xml += '<w:uiPriority w:val="9"/><w:unhideWhenUsed/><w:qFormat/>'
        elif sid.startswith('TOC'):
            xml += '<w:uiPriority w:val="39"/><w:unhideWhenUsed/>'
        else:
            xml += '<w:qFormat/>'
        xml += body + '</w:style>'
        el = frag(xml)
        if sid in have:
            root.replace(have[sid], el)
        else:
            root.append(el)
        have[sid] = el
    if extra_from is not None:
        other = etree.fromstring(extra_from)
        for s in other.findall(w('w:style')):
            sid = s.get(w('w:styleId'))
            if sid not in have:
                root.append(copy.deepcopy(s))
                have[sid] = s
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def tune_numbering(numbering_xml):
    """Make heading levels 2 and 3 of the template list match the new Heading 2/3 styles."""
    root = etree.fromstring(numbering_xml)
    for an in root.findall(w('w:abstractNum')):
        if an.get(w('w:abstractNumId')) != '0':
            continue
        for lvl in an.findall(w('w:lvl')):
            il = lvl.get(w('w:ilvl'))
            if il not in ('1', '2'):
                continue
            ind = lvl.find('w:pPr/w:ind', NS)
            if ind is not None:
                val = '720' if il == '1' else '864'
                ind.set(w('w:left'), val)
                ind.set(w('w:hanging'), val)
            pstyle = lvl.find(w('w:pStyle'))
            if pstyle is None:
                ps = E('w:pStyle', {'w:val': 'Heading2' if il == '1' else 'Heading3'})
                # pStyle goes after start/numFmt/lvlRestart, before isLgl/suff/lvlText
                anchor = lvl.find(w('w:lvlText'))
                anchor.addprevious(ps)
            rpr = lvl.find(w('w:rPr'))
            if rpr is not None:
                for t in ('w:sz', 'w:szCs'):
                    el = rpr.find(w(t))
                    if el is None:
                        el = E(t)
                        rpr.append(el)
                    el.set(w('w:val'), '26' if il == '1' else '23')
                col = rpr.find(w('w:color'))
                if col is not None:
                    col.set(w('w:val'), HEAD_BLUE)
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


# ------------------------------------------------------------- package io
class Package:
    def __init__(self, path):
        with zipfile.ZipFile(path) as z:
            self.parts = {n: z.read(n) for n in z.namelist()}
            self.order = z.namelist()

    def xml(self, name):
        return etree.fromstring(self.parts[name])

    def set_xml(self, name, root):
        self.parts[name] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        if name not in self.order:
            self.order.append(name)

    def set_bytes(self, name, data):
        self.parts[name] = data
        if name not in self.order:
            self.order.append(name)

    def delete(self, name):
        self.parts.pop(name, None)
        if name in self.order:
            self.order.remove(name)

    def save(self, path):
        tmp = path + '.tmp'
        with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
            names = [n for n in self.order if n in self.parts]
            if '[Content_Types].xml' in names:
                names.remove('[Content_Types].xml')
                names.insert(0, '[Content_Types].xml')
            for n in names:
                z.writestr(n, self.parts[n])
        os.replace(tmp, path)


def make_reference_doc(template, out):
    pkg = Package(template)
    pkg.parts['word/styles.xml'] = install_styles(pkg.parts['word/styles.xml'])
    pkg.parts['word/numbering.xml'] = tune_numbering(pkg.parts['word/numbering.xml'])
    pkg.save(out)


# --------------------------------------------------------- body processing
def pstyle(p):
    ps = p.find('w:pPr/w:pStyle', NS)
    return ps.get(w('w:val')) if ps is not None else None


def ensure_ppr(p):
    ppr = p.find(w('w:pPr'))
    if ppr is None:
        ppr = E('w:pPr')
        p.insert(0, ppr)
    return ppr


PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr',
             'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
             'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
             'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc',
             'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
             'sectPr', 'pPrChange']


def set_ppr_child(ppr, name, el):
    """Insert/replace a pPr child respecting the schema order."""
    old = ppr.find(w('w:' + name))
    if old is not None:
        ppr.replace(old, el)
        return el
    pos = PPR_ORDER.index(name)
    for i, child in enumerate(ppr):
        cname = etree.QName(child).localname
        if cname in PPR_ORDER and PPR_ORDER.index(cname) > pos:
            ppr.insert(i, el)
            return el
    ppr.append(el)
    return el


def fix_ppr_order(ppr):
    kids = list(ppr)
    def key(c):
        n = etree.QName(c).localname
        return PPR_ORDER.index(n) if n in PPR_ORDER else len(PPR_ORDER)
    kids_sorted = sorted(kids, key=key)
    if kids_sorted != kids:
        for c in kids:
            ppr.remove(c)
        for c in kids_sorted:
            ppr.append(c)


RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline',
             'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing',
             'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText',
             'vertAlign', 'rtl', 'cs', 'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath']
TRPR_ORDER = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit', 'trHeight',
              'tblHeader', 'tblCellSpacing', 'jc', 'hidden']


def fix_schema_order(root):
    """Sort w:rPr / w:trPr children into schema order (Word rejects out-of-order children)."""
    for tag, order in (('w:rPr', RPR_ORDER), ('w:trPr', TRPR_ORDER)):
        for el in root.iter(w(tag)):
            kids = list(el)
            def key(c):
                q = etree.QName(c)
                if q.namespace != W_NS:
                    return len(order) + 1
                return order.index(q.localname) if q.localname in order else len(order)
            srt = sorted(kids, key=key)
            if srt != kids:
                for c in kids:
                    el.remove(c)
                for c in srt:
                    el.append(c)
    for ppr in root.iter(w('w:pPr')):
        fix_ppr_order(ppr)


ID_CELL_RE = re.compile(r'^[A-Z][A-Za-z0-9]{0,6}(?:[-_.][A-Za-z0-9]{1,6}){1,3}[a-z]?$')


def cell_text(tc):
    return ' '.join(text_of(p) for p in tc.findall('.//w:p', NS)).strip()


def compute_widths(rows_text, total):
    ncol = max(len(r) for r in rows_text)
    cw, pad = 100, 200
    minw, desired = [], []
    for c in range(ncol):
        col = [r[c] if c < len(r) else '' for r in rows_text]
        tokens = [t for txt in col for t in re.split(r'\s+', txt) if t]
        longest = max([len(t) for t in tokens] + [3])
        head_tokens = [t for t in re.split(r'\s+', col[0]) if t] if col else []
        head_long = max([len(t) for t in head_tokens] + [0])
        mn = max(min(longest, 28) * 118, head_long * 122) + 240
        mx = max([len(t) for t in col] + [3])
        avg = sum(len(t) for t in col[1:]) / max(1, len(col) - 1)
        des = (0.5 * min(mx, 400) + 0.5 * min(avg * 1.6, 400)) * cw + pad
        minw.append(min(mn, int(total * 0.4)))
        desired.append(max(des, minw[-1]))
    widths = list(minw)
    if sum(widths) > total:
        f = total / sum(widths)
        widths = [x * f for x in widths]
    else:
        extra = total - sum(widths)
        need = [d - m for d, m in zip(desired, widths)]
        tn = sum(need)
        if tn > 0 and extra <= tn:
            widths = [m + n * extra / tn for m, n in zip(widths, need)]
        else:
            widths = [m + n for m, n in zip(widths, need)]
            left = total - sum(widths)
            s = sum(widths)
            widths = [x + left * x / s for x in widths]
    widths = [int(round(x)) for x in widths]
    widths[-1] += total - sum(widths)
    return widths


def table_rows_text(tbl):
    return [[cell_text(tc) for tc in tr.findall(w('w:tc'))] for tr in tbl.findall(w('w:tr'))]


def format_table(tbl, total_w, header_fill=HEAD_BLUE, zebra=True, restyle_text=True, widths=None):
    rows = tbl.findall(w('w:tr'))
    if not rows:
        return
    rows_text = table_rows_text(tbl)
    ncol = max(len(r) for r in rows_text)
    if widths is None or len(widths) != ncol:
        widths = compute_widths(rows_text, total_w)
    # tblPr
    old = tbl.find(w('w:tblPr'))
    tblpr = frag(
        '<w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="5000" w:type="pct"/><w:tblInd w:w="0" w:type="dxa"/>'
        '<w:tblBorders>' + ''.join('<w:%s w:val="single" w:sz="4" w:space="0" w:color="%s"/>' % (s, GRID)
                                   for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')) +
        '</w:tblBorders><w:tblLayout w:type="fixed"/>'
        '<w:tblCellMar><w:top w:w="45" w:type="dxa"/><w:left w:w="85" w:type="dxa"/>'
        '<w:bottom w:w="45" w:type="dxa"/><w:right w:w="85" w:type="dxa"/></w:tblCellMar>'
        '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'
        '</w:tblPr>')
    if old is not None:
        tbl.replace(old, tblpr)
    else:
        tbl.insert(0, tblpr)
    grid = tbl.find(w('w:tblGrid'))
    newgrid = E('w:tblGrid')
    for x in widths:
        newgrid.append(E('w:gridCol', {'w:w': x}))
    if grid is not None:
        tbl.replace(grid, newgrid)
    else:
        tblpr.addnext(newgrid)
    for ri, tr in enumerate(rows):
        header = ri == 0
        trpr = tr.find(w('w:trPr'))
        if trpr is None:
            trpr = E('w:trPr')
            tr.insert(0, trpr)
        for ch in list(trpr):
            if etree.QName(ch).localname in ('tblHeader', 'cantSplit'):
                trpr.remove(ch)
        total_chars = sum(len(t) for t in rows_text[ri])
        if total_chars < 700:
            trpr.append(E('w:cantSplit'))
        if header:
            trpr.append(E('w:tblHeader'))
        for ci, tc in enumerate(tr.findall(w('w:tc'))):
            tcpr = tc.find(w('w:tcPr'))
            if tcpr is None:
                tcpr = E('w:tcPr')
                tc.insert(0, tcpr)
            span = tcpr.find(w('w:gridSpan'))
            sp = int(wattr(span, 'w:val', '1')) if span is not None else 1
            for ch in list(tcpr):
                if etree.QName(ch).localname in ('tcW', 'shd', 'noWrap', 'vAlign'):
                    tcpr.remove(ch)
            tcpr.insert(0, E('w:tcW', {'w:w': sum(widths[ci:ci + sp]), 'w:type': 'dxa'}))
            txt = rows_text[ri][ci] if ci < len(rows_text[ri]) else ''
            if header:
                tcpr.append(E('w:shd', {'w:val': 'clear', 'w:color': 'auto', 'w:fill': header_fill}))
            elif zebra and ri % 2 == 0:
                tcpr.append(E('w:shd', {'w:val': 'clear', 'w:color': 'auto', 'w:fill': BAND}))
            if not header and ID_CELL_RE.match(txt):
                tcpr.append(E('w:noWrap'))
            tcpr.append(E('w:vAlign', {'w:val': 'center' if header else 'top'}))
            # order inside tcPr: tcW, gridSpan, vMerge, tcBorders, shd, noWrap, tcMar, ..., vAlign
            order = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
                     'textDirection', 'tcFitText', 'vAlign', 'hideMark']
            kids = sorted(list(tcpr), key=lambda c: order.index(etree.QName(c).localname)
                          if etree.QName(c).localname in order else 99)
            for c in list(tcpr):
                tcpr.remove(c)
            for c in kids:
                tcpr.append(c)
            if not restyle_text:
                if header:
                    for r in tc.iter(w('w:r')):
                        rpr = r.find(w('w:rPr'))
                        if rpr is None:
                            rpr = E('w:rPr')
                            r.insert(0, rpr)
                        col = rpr.find(w('w:color'))
                        if col is None:
                            col = E('w:color')
                            rpr.append(col)
                        col.set(w('w:val'), 'FFFFFF')
                        for att in list(col.attrib):
                            if 'theme' in att:
                                del col.attrib[att]
                continue
            for p in tc.findall(w('w:p')):
                ppr = ensure_ppr(p)
                st = pstyle(p)
                if st in (None, 'Compact', 'BodyText', 'FirstParagraph', 'Normal'):
                    set_ppr_child(ppr, 'pStyle', E('w:pStyle', {'w:val': 'TableHeading' if header else 'TableText'}))
                if p.find('w:pPr/w:numPr', NS) is not None:
                    set_ppr_child(ppr, 'spacing', E('w:spacing', {'w:before': '0', 'w:after': '20'}))
                if header:
                    for r in p.iter(w('w:r')):
                        rpr = r.find(w('w:rPr'))
                        if rpr is None:
                            rpr = E('w:rPr')
                            r.insert(0, rpr)
                        if rpr.find(w('w:color')) is None:
                            rpr.append(E('w:color', {'w:val': 'FFFFFF'}))
                        else:
                            rpr.find(w('w:color')).set(w('w:val'), 'FFFFFF')
    return ncol


def normalise_jc(root):
    for jc in root.iter(w('w:jc')):
        v = jc.get(w('w:val'))
        if v == 'start':
            jc.set(w('w:val'), 'left')
        elif v == 'end':
            jc.set(w('w:val'), 'right')


def heading_level(p):
    st = pstyle(p) or ''
    m = re.match(r'^Heading([1-9])$', st)
    return int(m.group(1)) if m else 0


def plain_heading(p, level):
    """Reset heading runs to plain text (headings take their look from the style)."""
    txt = text_of(p).strip()
    ppr = p.find(w('w:pPr'))
    for ch in list(p):
        if ch is not ppr:
            p.remove(ch)
    if level >= 4:
        set_ppr_child(ppr, 'pStyle', E('w:pStyle', {'w:val': 'Heading4'}))
    r = E('w:r')
    t = E('w:t')
    t.text = txt
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    p.append(r)
    return txt


def widen_num_tab(ppr):
    """Room for two-digit heading numbers (the template tab at 360 twips only fits '9.')."""
    tabs = ppr.find(w('w:tabs'))
    if tabs is not None:
        for t in tabs.findall(w('w:tab')):
            if t.get(w('w:val')) == 'num':
                t.set(w('w:pos'), '600')


def make_h1(proto, text):
    p = copy.deepcopy(proto)
    ppr = p.find(w('w:pPr'))
    run_rpr = None
    for r in p.findall(w('w:r')):
        if r.find(w('w:t')) is not None and text_of(r).strip():
            run_rpr = copy.deepcopy(r.find(w('w:rPr')))
            break
    for ch in list(p):
        if ch is not ppr:
            p.remove(ch)
    numpr = ppr.find(w('w:numPr'))
    if numpr is not None:
        numpr.find(w('w:numId')).set(w('w:val'), '1')
    set_ppr_child(ppr, 'spacing', E('w:spacing', {'w:before': '360', 'w:after': '160'}))
    widen_num_tab(ppr)
    r = E('w:r')
    if run_rpr is not None:
        r.append(run_rpr)
    t = E('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    p.append(r)
    return p


def empty_para(spacing_after=None):
    p = E('w:p')
    if spacing_after is not None:
        ppr = E('w:pPr')
        ppr.append(E('w:spacing', {'w:before': '0', 'w:after': str(spacing_after)}))
        p.append(ppr)
    return p


def is_empty_para(el):
    return (el.tag == w('w:p') and not text_of(el).strip() and el.find('.//w:drawing', NS) is None
            and el.find('.//w:pict', NS) is None and el.find('.//w:sectPr', NS) is None
            and el.find('.//w:br[@w:type="page"]', NS) is None and el.find('.//w:fldChar', NS) is None)


def plan_landscape(elems):
    """Return list of (start, end) index ranges (inclusive) that go into landscape sections."""
    wide = []
    for i, el in enumerate(elems):
        if el.tag == w('w:tbl'):
            grid = el.find(w('w:tblGrid'))
            ncol = len(grid.findall(w('w:gridCol'))) if grid is not None else 0
            if ncol >= 6:
                wide.append(i)
        elif el.tag == w('w:p') and pstyle(el) == 'FigureWide':
            wide.append(i)
    regions = []
    for i in wide:
        start = i
        j = i - 1
        steps = 0
        # pull in captions, a short intro and the heading of the section
        while j >= 0 and steps < 4:
            el = elems[j]
            if el.tag != w('w:p'):
                break
            st = pstyle(el)
            if el.find('.//w:sectPr', NS) is not None:
                break
            start = j
            if heading_level(el):
                # keep consecutive headings together (e.g. H1 directly followed by H2)
                while j - 1 >= 0 and elems[j - 1].tag == w('w:p') and heading_level(elems[j - 1]):
                    j -= 1
                    start = j
                break
            if st not in ('TableCaption', 'Caption') and len(text_of(el)) > 400:
                start = j + 1
                break
            j -= 1
            steps += 1
        else:
            # no heading found within reach: only keep captions directly above
            start = i
            j = i - 1
            while j >= 0 and elems[j].tag == w('w:p') and pstyle(elems[j]) in ('TableCaption',):
                start = j
                j -= 1
        end = i
        if pstyle(elems[i]) == 'FigureWide' and i + 1 < len(elems) and pstyle(elems[i + 1]) == 'Caption':
            end = i + 1
        if regions and start <= regions[-1][1] + 1:
            regions[-1] = (regions[-1][0], end)
        else:
            regions.append((start, end))
    # merge regions separated by a small amount of non-table content
    merged = []
    for r in regions:
        if merged:
            gap = elems[merged[-1][1] + 1:r[0]]
            if len(gap) <= 3 and all(g.tag == w('w:p') and g.find('.//w:drawing', NS) is None for g in gap):
                merged[-1] = (merged[-1][0], r[1])
                continue
        merged.append(r)
    return merged


# --------------------------------------------------------- frame filling
def strip_rtl(rpr):
    if rpr is None:
        return None
    rpr = copy.deepcopy(rpr)
    for t in ('w:rtl', 'w:lang'):
        for el in rpr.findall(w(t)):
            rpr.remove(el)
    rf = rpr.find(w('w:rFonts'))
    if rf is not None and rf.get(w('w:hint')) == 'cs':
        del rf.attrib[w('w:hint')]
    return rpr


def fill_cell(tc, lines, align=None):
    """Replace the content of a template cell, keeping its paragraph and run look."""
    if isinstance(lines, str):
        lines = [lines]
    paras = tc.findall(w('w:p'))
    p = paras[0]
    for extra in paras[1:]:
        tc.remove(extra)
    ppr = p.find(w('w:pPr'))
    rpr = None
    for r in p.findall('.//w:r', NS):
        if r.find(w('w:t')) is not None:
            rpr = strip_rtl(r.find(w('w:rPr')))
            break
    if rpr is None and ppr is not None and ppr.find(w('w:rPr')) is not None:
        rpr = strip_rtl(ppr.find(w('w:rPr')))
    if ppr is not None and ppr.find(w('w:rPr')) is not None:
        ppr.replace(ppr.find(w('w:rPr')), strip_rtl(ppr.find(w('w:rPr'))))
    for ch in list(p):
        if ch is not ppr:
            p.remove(ch)
    if align and ppr is not None:
        set_ppr_child(ppr, 'jc', E('w:jc', {'w:val': align}))
    for k, line in enumerate(lines):
        r = E('w:r')
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        if k:
            r.append(E('w:br'))
        t = E('w:t')
        t.text = line
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t)
        p.append(r)


def set_rows(tbl, data, proto_index=1, align=None, min_height=None):
    rows = tbl.findall(w('w:tr'))
    proto = copy.deepcopy(rows[proto_index])
    for r in rows[1:]:
        tbl.remove(r)
    for vals in data:
        tr = copy.deepcopy(proto)
        if min_height:
            trpr = tr.find(w('w:trPr'))
            h = trpr.find(w('w:trHeight')) if trpr is not None else None
            if h is not None:
                h.set(w('w:val'), str(min_height))
        for k, tc in enumerate(tr.findall(w('w:tc'))):
            v = vals[k] if k < len(vals) else ''
            fill_cell(tc, v, align=(align[k] if align else None))
        tbl.append(tr)


def set_textbox(anchor, segments, jc=None, size=None):
    """segments: list of (text, bold). Rewrites the first paragraph of every txbxContent in the anchor."""
    for tx in anchor.iter(w('w:txbxContent')):
        paras = tx.findall(w('w:p'))
        p = paras[0]
        for extra in paras[1:]:
            tx.remove(extra)
        runs = [r for r in p.findall(w('w:r')) if r.find(w('w:t')) is not None]
        reg = bold = None
        for r in runs:
            rpr = r.find(w('w:rPr'))
            is_b = rpr is not None and rpr.find(w('w:b')) is not None
            if is_b and bold is None:
                bold = copy.deepcopy(rpr)
            if not is_b and reg is None:
                reg = copy.deepcopy(rpr)
        if bold is None and reg is not None:
            bold = copy.deepcopy(reg)
            bold.insert(0, E('w:b'))
        if reg is None and bold is not None:
            reg = copy.deepcopy(bold)
            for t in ('w:b', 'w:bCs'):
                for el in reg.findall(w(t)):
                    reg.remove(el)
        ppr = p.find(w('w:pPr'))
        for ch in list(p):
            if ch is not ppr:
                p.remove(ch)
        if ppr is None:
            ppr = E('w:pPr')
            p.insert(0, ppr)
        if jc:
            set_ppr_child(ppr, 'jc', E('w:jc', {'w:val': jc}))
        for text, is_bold in segments:
            r = E('w:r')
            rpr = copy.deepcopy(bold if is_bold else reg)
            if rpr is not None:
                if size:
                    for t in ('w:sz', 'w:szCs'):
                        el = rpr.find(w(t))
                        if el is None:
                            el = E(t)
                            rpr.append(el)
                        el.set(w('w:val'), str(size))
                r.append(rpr)
            t = E('w:t')
            t.text = text
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            r.append(t)
            p.append(r)
        if ppr.find(w('w:rPr')) is not None and size:
            prp = ppr.find(w('w:rPr'))
            for t in ('w:sz', 'w:szCs'):
                el = prp.find(w(t))
                if el is not None:
                    el.set(w('w:val'), str(size))


def set_anchor_geometry(anchor, x=None, y=None, cx=None, cy=None, rel_h=None, align_h=None):
    ph = anchor.find(w('wp:positionH'))
    if align_h is not None:
        for ch in list(ph):
            ph.remove(ch)
        ph.set('relativeFrom', rel_h or 'margin')
        el = etree.SubElement(ph, w('wp:align'))
        el.text = align_h
    elif x is not None:
        for ch in list(ph):
            ph.remove(ch)
        ph.set('relativeFrom', rel_h or 'page')
        el = etree.SubElement(ph, w('wp:posOffset'))
        el.text = str(int(x))
    if y is not None:
        pv = anchor.find(w('wp:positionV'))
        off = pv.find(w('wp:posOffset'))
        off.text = str(int(y))
    ext = anchor.find(w('wp:extent'))
    if cx is not None:
        ext.set('cx', str(int(cx)))
    if cy is not None:
        ext.set('cy', str(int(cy)))
    for x_ in anchor.iter(w('a:ext')):
        if x_.getparent().tag == w('a:xfrm'):
            if cx is not None:
                x_.set('cx', str(int(cx)))
            if cy is not None:
                x_.set('cy', str(int(cy)))


def estimate_lines(text, pt, width_cm, bold=True):
    char_cm = pt * (0.60 if bold else 0.55) * 2.54 / 72
    per_line = max(1, int(width_cm / char_cm))
    words = text.split()
    lines, cur = 1, 0
    for wd in words:
        add = len(wd) + (1 if cur else 0)
        if cur + add > per_line and cur:
            lines += 1
            cur = len(wd)
        else:
            cur += add
    return lines


def fill_cover(cover_p, meta):
    # remove VML fallbacks so only the DrawingML text boxes need to be maintained
    for fb in cover_p.findall('.//mc:Fallback', NS):
        fb.getparent().remove(fb)
    anchors = [a for a in cover_p.iter(w('wp:anchor')) if a.find('.//w:txbxContent', NS) is not None]
    by = {}
    for a in anchors:
        t = text_of(a)
        if t.startswith('Business Requirement'):
            by['title'] = a
        elif t.startswith('Sub'):
            by['subtitle'] = a
        elif t.startswith('Prepared'):
            by['prepared'] = a
        elif t.startswith('Version'):
            by['version'] = a
        elif t.startswith('Date'):
            by['date'] = a
    page_w_emu = PAGE_W * 635
    box_w_cm = 16.0
    title = meta['title'] or 'Untitled'
    t_lines = estimate_lines(title, 22, box_w_cm - 0.5)
    line_emu = int(22 * 1.32 * 12700)
    if 'title' in by:
        set_textbox(by['title'], [(title, True)], jc='center')
        set_anchor_geometry(by['title'], cx=box_w_cm * EMU_PER_CM, cy=max(476885, t_lines * line_emu + 91440),
                            align_h='center', rel_h='margin')
    sub = meta['subtitle']
    if meta['id']:
        sub = (sub + ' \u00b7 ' + meta['id']) if sub else meta['id']
    if 'subtitle' in by:
        size = 36
        for pt in (18, 16, 14):
            size = pt * 2
            if estimate_lines(sub, pt, box_w_cm - 0.5) == 1:
                break
        s_lines = estimate_lines(sub, size / 2, box_w_cm - 0.5)
        y0 = 1979157 + (t_lines - 1) * line_emu
        set_textbox(by['subtitle'], [(sub, True)], jc='center', size=size)
        set_anchor_geometry(by['subtitle'], y=y0, cx=box_w_cm * EMU_PER_CM,
                            cy=max(337820, int(s_lines * size / 2 * 1.32 * 12700) + 91440),
                            align_h='center', rel_h='margin')
    org = split_org(meta['prepared_by'])[0] or 'iorta TechNXT'
    if 'prepared' in by:
        set_textbox(by['prepared'], [('Prepared by \u2013 ', False), (org, True)])
        set_anchor_geometry(by['prepared'], x=194310, cx=3000000, rel_h='page')
    if 'version' in by:
        vw = 1700000
        set_textbox(by['version'], [('Version \u2013 ', False), (meta['version'], True)], jc='center')
        set_anchor_geometry(by['version'], x=page_w_emu / 2 - vw / 2, cx=vw, rel_h='page')
    if 'date' in by:
        right = 5762625 + 1628775
        dw = 2600000
        set_textbox(by['date'], [('Date \u2013 ', False), (meta['date_long'], True)], jc='right')
        set_anchor_geometry(by['date'], x=right - dw, cx=dw, rel_h='page')


# ------------------------------------------------------------------- TOC
def toc_entry(level, label, text, bookmark, page, first=False, last=False):
    rpr = '<w:rPr><w:noProof/></w:rPr>'
    parts = ['<w:p><w:pPr><w:pStyle w:val="TOC%d"/></w:pPr>' % level]
    if first:
        parts.append('<w:r>%s<w:fldChar w:fldCharType="begin"/></w:r>'
                     '<w:r>%s<w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
                     '<w:r>%s<w:fldChar w:fldCharType="separate"/></w:r>' % (rpr, rpr, rpr))
    esc = lambda s: s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    parts.append('<w:hyperlink w:anchor="%s" w:history="1">' % bookmark)
    if label:
        parts.append('<w:r>%s<w:t>%s</w:t></w:r><w:r>%s<w:tab/></w:r>' % (rpr, esc(label), rpr))
    parts.append('<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r>' % (rpr, esc(text)))
    wh = '<w:rPr><w:noProof/><w:webHidden/></w:rPr>'
    parts.append('<w:r>%s<w:tab/></w:r><w:r>%s<w:fldChar w:fldCharType="begin"/></w:r>'
                 '<w:r>%s<w:instrText xml:space="preserve"> PAGEREF %s \\h </w:instrText></w:r>'
                 '<w:r>%s<w:fldChar w:fldCharType="separate"/></w:r><w:r>%s<w:t>%s</w:t></w:r>'
                 '<w:r>%s<w:fldChar w:fldCharType="end"/></w:r>' % (wh, wh, wh, bookmark, wh, wh, page, wh))
    parts.append('</w:hyperlink>')
    if last:
        parts.append('<w:r>%s<w:fldChar w:fldCharType="end"/></w:r>' % rpr)
    parts.append('</w:p>')
    return frag(''.join(parts))


# ------------------------------------------------------------- the build
class Builder:
    def __init__(self, args):
        self.args = args
        self.quiet = args.quiet
        self.work = tempfile.mkdtemp(prefix='mdocx_')

    def run(self):
        a = self.args
        with open(a.input, encoding='utf-8') as f:
            text = f.read()
        fm, body = split_front_matter(text)
        if fm is None:
            raise SystemExit('ERROR: %s has no YAML front matter' % a.input)
        meta = normalise_meta(fm)
        if not meta['title']:
            raise SystemExit('ERROR: front matter has no title')
        self.meta = meta
        md_dir = os.path.dirname(os.path.abspath(a.input))
        body = normalise_markdown(body, meta)
        pre, nfig, ntab = preprocess(body, md_dir, a.img_dir, self.work, self.quiet)
        pre_path = os.path.join(self.work, 'body.md')
        with open(pre_path, 'w', encoding='utf-8') as f:
            f.write(pre)
        ref = os.path.join(self.work, 'reference.docx')
        make_reference_doc(a.template, ref)
        body_docx = os.path.join(self.work, 'body.docx')
        cmd = ['pandoc', pre_path, '-f', 'gfm+fenced_divs+attributes+footnotes+smart-yaml_metadata_block',
               '-t', 'docx', '--reference-doc', ref, '--lua-filter', LUA_FILTER,
               '--resource-path', md_dir, '-o', body_docx]
        res = subprocess.run(cmd, capture_output=True, text=True)
        if res.returncode != 0:
            raise SystemExit('pandoc failed: ' + res.stderr)
        for line in res.stderr.splitlines():
            if line.strip():
                warn('pandoc: ' + line.strip())
        self.assemble(body_docx, a.output)
        self.stats = {'figures': nfig, 'table_captions': ntab}
        toc_info = self.refresh_toc(a.output)
        report = self.validate(a.output)
        report.update(toc_info)
        report.update(self.stats)
        if a.preview:
            self.preview(a.output, a.preview)
        if not a.keep_work:
            shutil.rmtree(self.work, ignore_errors=True)
        report['warnings'] = list(WARNINGS)
        return report

    # ------------------------------------------------------------------
    def assemble(self, body_docx, out_path):
        a = self.args
        meta = self.meta
        master = Package(a.template)
        src = Package(body_docx)

        # styles / numbering
        master.parts['word/styles.xml'] = install_styles(master.parts['word/styles.xml'],
                                                         extra_from=src.parts['word/styles.xml'])
        num = etree.fromstring(tune_numbering(master.parts['word/numbering.xml']))
        snum = src.xml('word/numbering.xml')
        have_abs = {x.get(w('w:abstractNumId')) for x in num.findall(w('w:abstractNum'))}
        have_num = {x.get(w('w:numId')) for x in num.findall(w('w:num'))}
        first_num = num.find(w('w:num'))
        for an in snum.findall(w('w:abstractNum')):
            if an.get(w('w:abstractNumId')) not in have_abs:
                self.style_list_levels(an)
                first_num.addprevious(copy.deepcopy(an))
        for n in snum.findall(w('w:num')):
            if n.get(w('w:numId')) not in have_num:
                num.append(copy.deepcopy(n))
        master.set_xml('word/numbering.xml', num)

        doc = master.xml('word/document.xml')
        body = doc.find(w('w:body'))
        sbody = src.xml('word/document.xml').find(w('w:body'))
        rels = master.xml('word/_rels/document.xml.rels')
        srels = {r.get('Id'): r for r in src.xml('word/_rels/document.xml.rels')}
        ct = master.xml('[Content_Types].xml')

        # ---------- locate template frame
        kids = list(body)
        final_sectpr = kids[-1]
        assert final_sectpr.tag == w('w:sectPr')

        def find_h1(txt):
            for el in body:
                if el.tag == w('w:p') and pstyle(el) == 'Heading1' and text_of(el).strip().lower() == txt.lower():
                    return el
            raise SystemExit('template heading not found: ' + txt)

        h_dc = find_h1('Document Control')
        h_acr = find_h1('Acronyms')
        h_intro = find_h1('Introduction')
        h_sign = find_h1('Requirements Sign-Off')
        h1_proto = copy.deepcopy(h_acr)
        tables = body.findall(w('w:tbl'))
        t_dc, t_acr, t_so, t_sig = tables[:4]

        # ---------- cover
        cover_p = next(el for el in body if el.find('.//w:txbxContent', NS) is not None)
        fill_cover(cover_p, meta)

        # ---------- TOC: remove cached entries and padding paragraphs up to Document Control
        toc_head = next(el for el in body if el.tag == w('w:p') and pstyle(el) == 'TOCHeading')
        el = toc_head.getnext()
        while el is not None and el is not h_dc:
            nxt = el.getnext()
            body.remove(el)
            el = nxt
        toc_anchor = toc_head  # entries are inserted after this later

        # ---------- Document Control heading: drop stray TOC field end, start a new page
        for r in h_dc.findall(w('w:r')):
            if r.find(w('w:fldChar')) is not None:
                h_dc.remove(r)
        for h in (h_dc, h_acr):
            hp = h.find(w('w:pPr'))
            set_ppr_child(hp, 'spacing', E('w:spacing', {'w:before': '360', 'w:after': '160'}))
            widen_num_tab(hp)
        set_ppr_child(h_dc.find(w('w:pPr')), 'pageBreakBefore', E('w:pageBreakBefore'))

        # ---------- Document Control table
        def two_lines(s):
            org, unit = split_org(s)
            return [org, unit] if unit else [org]
        set_rows(t_dc, [[meta['date_short'], meta['version'], two_lines(meta['prepared_by']),
                         two_lines(meta['reviewed_by']), two_lines(meta['approved_by']),
                         meta['change_history']]])
        # ---------- Acronyms
        acr = meta['acronyms'] or [('\u2013', 'No acronyms used in this document')]
        set_rows(t_acr, [[a_, m_] for a_, m_ in acr])
        # ---------- clean padding between the frame tables
        self.squeeze_empties(body, t_dc, h_acr, keep=1)
        self.squeeze_empties(body, t_acr, h_intro, keep=0)

        # ---------- Sign-off
        so_title = 'Requirements Sign-Off' if meta['id'] in REQUIREMENT_DOC_IDS else 'Sign-Off'
        new_sign = make_h1(h1_proto, so_title)
        set_ppr_child(new_sign.find(w('w:pPr')), 'pageBreakBefore', E('w:pageBreakBefore'))
        h_sign.addprevious(new_sign)
        body.remove(h_sign)
        h_sign = new_sign
        so = meta['signoff'] or [('No open items', '', 'Closed')]
        set_rows(t_so, [[str(i + 1), d, o, s] for i, (d, o, s) in enumerate(so)],
                 align=['center', 'left', 'center', 'center'])
        sig_rows = []
        for who in (meta['approved_by'], meta['reviewed_by'], meta['prepared_by']):
            if who:
                org, unit = split_org(who)
                sig_rows.append(['', '', org, unit])
        if not sig_rows:
            sig_rows = [['', '', 'TASCO Insurance', ''], ['', '', 'iorta TechNXT', '']]
        set_rows(t_sig, sig_rows, min_height=760)
        self.squeeze_empties(body, h_sign, t_so, keep=0)
        self.squeeze_empties(body, t_so, t_sig, keep=1)
        # back cover: exactly one page break before it
        el = t_sig.getnext()
        breaks = []
        while el is not None and el.tag != w('w:sectPr'):
            if el.tag == w('w:p') and el.find('.//w:br[@w:type="page"]', NS) is not None:
                breaks.append(el)
            el = el.getnext()
        for b in breaks[1:]:
            body.remove(b)
        for tbl in (t_dc, t_acr, t_so, t_sig):
            format_table(tbl, PORTRAIT_TW, zebra=False, restyle_text=False)
            # template frame tables keep their own widths
            tbl.find('w:tblPr/w:tblLayout', NS).set(w('w:type'), 'autofit')
        # restore template width of the acronyms table (70 %)
        t_acr.find('w:tblPr/w:tblW', NS).set(w('w:w'), '3480')

        # ---------- remove template placeholder sections
        el = h_intro
        while el is not None and el is not h_sign:
            nxt = el.getnext()
            body.remove(el)
            el = nxt

        # ---------- convert pandoc body
        elems = [c for c in sbody if c.tag != w('w:sectPr')]
        out_elems = []
        prev_table = False
        for el in elems:
            for bm in el.xpath('.//w:bookmarkStart|.//w:bookmarkEnd', namespaces=NS):
                bm.getparent().remove(bm)
            if el.tag in (w('w:bookmarkStart'), w('w:bookmarkEnd')):
                continue
            normalise_jc(el)
            if el.tag == w('w:p'):
                lvl = heading_level(el)
                if lvl == 1:
                    el = make_h1(h1_proto, text_of(el).strip())
                elif lvl >= 2:
                    plain_heading(el, lvl)
                elif prev_table:
                    ppr = ensure_ppr(el)
                    sp = ppr.find(w('w:spacing'))
                    if sp is None:
                        sp = set_ppr_child(ppr, 'spacing', E('w:spacing'))
                    sp.set(w('w:before'), '160')
                prev_table = False
            elif el.tag == w('w:tbl'):
                if prev_table:
                    out_elems.append(empty_para(0))
                prev_table = True
            out_elems.append(el)

        # first body element starts on a new page
        if out_elems:
            first = out_elems[0]
            if first.tag == w('w:p'):
                set_ppr_child(ensure_ppr(first), 'pageBreakBefore', E('w:pageBreakBefore'))
            else:
                pb = E('w:p')
                ppr = E('w:pPr')
                ppr.append(E('w:pageBreakBefore'))
                pb.append(ppr)
                out_elems.insert(0, pb)

        # landscape planning (before table widths are fixed)
        regions = plan_landscape(out_elems)
        in_ls = set()
        for s, e in regions:
            in_ls.update(range(s, e + 1))
        # tables with the same header and orientation share one set of column widths
        groups = {}
        for i, el in enumerate(out_elems):
            if el.tag == w('w:tbl'):
                rt = table_rows_text(el)
                if rt:
                    key = (tuple(rt[0]), i in in_ls)
                    if key in groups:
                        groups[key].extend(rt[1:])
                    else:
                        groups[key] = list(rt)
        group_widths = {k: compute_widths(v, LANDSCAPE_TW if k[1] else PORTRAIT_TW) for k, v in groups.items()}
        for i, el in enumerate(out_elems):
            if el.tag == w('w:tbl'):
                rt = table_rows_text(el)
                key = (tuple(rt[0]), i in in_ls) if rt else None
                format_table(el, LANDSCAPE_TW if i in in_ls else PORTRAIT_TW, widths=group_widths.get(key))
            elif el.tag == w('w:p') and pstyle(el) == 'FigureWide':
                set_ppr_child(el.find(w('w:pPr')), 'pStyle', E('w:pStyle', {'w:val': 'Figure'}))

        # landscape footer part
        ls_footer_rid = None
        if regions:
            ls_footer_rid = self.add_landscape_footer(master, rels, ct)
        portrait_sect = copy.deepcopy(final_sectpr)
        for k in [k for k in portrait_sect.attrib]:
            del portrait_sect.attrib[k]
        land_sect = copy.deepcopy(portrait_sect)
        pg = land_sect.find(w('w:pgSz'))
        pg.set(w('w:w'), str(PAGE_H))
        pg.set(w('w:h'), str(PAGE_W))
        pg.set(w('w:orient'), 'landscape')
        for ref in land_sect.findall(w('w:footerReference')):
            if ref.get(w('w:type')) == 'default' and ls_footer_rid:
                ref.set(w('r:id'), ls_footer_rid)

        def sect_para(sect):
            p = E('w:p')
            ppr = E('w:pPr')
            ppr.append(E('w:spacing', {'w:before': '0', 'w:after': '0', 'w:line': '20', 'w:lineRule': 'exact'}))
            rpr = E('w:rPr')
            rpr.append(E('w:sz', {'w:val': '2'}))
            ppr.append(rpr)
            ppr.append(copy.deepcopy(sect))
            p.append(ppr)
            return p

        final_elems = []
        starts = {s: e for s, e in regions}
        ends = {e for s, e in regions}
        for i, el in enumerate(out_elems):
            if i in starts:
                final_elems.append(sect_para(portrait_sect))
                ppr = el.find(w('w:pPr')) if el.tag == w('w:p') else None
                if ppr is not None and ppr.find(w('w:pageBreakBefore')) is not None:
                    ppr.remove(ppr.find(w('w:pageBreakBefore')))
            final_elems.append(el)
            if i in ends:
                final_elems.append(sect_para(land_sect))
                nxt = out_elems[i + 1] if i + 1 < len(out_elems) else h_sign
                if nxt.tag == w('w:p'):
                    nppr = nxt.find(w('w:pPr'))
                    if nppr is not None and nppr.find(w('w:pageBreakBefore')) is not None:
                        nppr.remove(nppr.find(w('w:pageBreakBefore')))
        if regions and regions[0][0] == 0:
            # body begins in landscape: the portrait section break already starts a new page
            pass

        # ---------- relationships, media, footnotes
        self.rel_counter = 0
        media_n = [0]
        existing_ids = {r.get('Id') for r in rels}

        def new_rid():
            while True:
                self.rel_counter += 1
                rid = 'rIdG%d' % self.rel_counter
                if rid not in existing_ids:
                    existing_ids.add(rid)
                    return rid
        rid_map = {}

        def map_rel(old):
            if old in rid_map:
                return rid_map[old]
            r = srels.get(old)
            if r is None:
                warn('dangling relationship %s in pandoc output' % old)
                return old
            rid = new_rid()
            typ = r.get('Type')
            nr = etree.SubElement(rels, '{%s}Relationship' % NS['rel'])
            nr.set('Id', rid)
            nr.set('Type', typ)
            if r.get('TargetMode') == 'External':
                nr.set('Target', r.get('Target'))
                nr.set('TargetMode', 'External')
            else:
                target = r.get('Target')
                src_name = 'word/' + target.lstrip('/') if not target.startswith('word/') else target
                src_name = os.path.normpath(src_name).replace(os.sep, '/')
                data = src.parts.get(src_name)
                if data is None:
                    warn('missing part %s' % src_name)
                    return old
                media_n[0] += 1
                ext = os.path.splitext(src_name)[1].lower() or '.bin'
                name = 'media/body_%03d%s' % (media_n[0], ext)
                master.set_bytes('word/' + name, data)
                nr.set('Target', name)
                self.ensure_default_ct(ct, ext.lstrip('.'))
            rid_map[old] = rid
            return rid

        for el in final_elems:
            for node in el.iter():
                if node.tag in (w('w:headerReference'), w('w:footerReference')):
                    continue  # section breaks we created point at the template's own parts
                for att in list(node.attrib):
                    if att.startswith('{%s}' % R_NS):
                        node.set(att, map_rel(node.get(att)))
        self.merge_footnotes(master, src, final_elems)

        # ---------- splice body in place of the template placeholder sections
        for el in final_elems:
            h_sign.addprevious(el)

        # ---------- headings: bookmarks and TOC
        self.headings = self.collect_headings(body, toc_head)
        entries = []
        for k, hd in enumerate(self.headings):
            entries.append(toc_entry(hd['level'], hd['label'], hd['text'], hd['bookmark'], hd.get('page', ''),
                                     first=(k == 0), last=(k == len(self.headings) - 1)))
        anchor = toc_anchor
        for e in entries:
            anchor.addnext(e)
            anchor = e

        # ---------- unique drawing ids
        did = 5000
        for dp in body.iter(w('wp:docPr')):
            if dp.getparent() is not None and dp.get('id', '0').isdigit() and int(dp.get('id')) < 100000:
                did += 1
                dp.set('id', str(did))
        self.renumber_bookmarks(body)
        fix_schema_order(doc)

        master.set_xml('word/document.xml', doc)
        master.set_xml('word/_rels/document.xml.rels', rels)
        self.strip_custom_xml(master, rels, ct)
        master.set_xml('word/_rels/document.xml.rels', rels)
        master.set_xml('[Content_Types].xml', ct)
        self.set_settings(master)
        self.set_props(master)
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        master.save(out_path)
        self.master_path = out_path

    # ------------------------------------------------------------------
    @staticmethod
    def style_list_levels(an):
        """Pandoc list definitions: Segoe UI bullets in the heading blue, tidy indents."""
        for lvl in an.findall(w('w:lvl')):
            il = int(lvl.get(w('w:ilvl')))
            fmt = lvl.find(w('w:numFmt'))
            ppr = lvl.find(w('w:pPr'))
            if ppr is None:
                ppr = E('w:pPr')
                lvl.append(ppr)
            ind = ppr.find(w('w:ind'))
            if ind is None:
                ind = E('w:ind')
                ppr.append(ind)
            ind.set(w('w:left'), str(360 + il * 360))
            ind.set(w('w:hanging'), '284' if fmt is not None and fmt.get(w('w:val')) == 'bullet' else '360')
            if fmt is not None and fmt.get(w('w:val')) == 'bullet':
                txt = lvl.find(w('w:lvlText'))
                txt.set(w('w:val'), ['\u2022', '\u2013', '\u25e6'][il % 3])
                rpr = lvl.find(w('w:rPr'))
                if rpr is None:
                    rpr = E('w:rPr')
                    lvl.append(rpr)
                for ch in list(rpr):
                    rpr.remove(ch)
                rpr.append(frag(SEGOE))
                rpr.append(E('w:color', {'w:val': PRIMARY}))

    @staticmethod
    def squeeze_empties(body, after, before, keep=0):
        el = after.getnext()
        empties = []
        while el is not None and el is not before:
            if is_empty_para(el):
                empties.append(el)
            el = el.getnext()
        for e in empties[keep:]:
            body.remove(e)

    @staticmethod
    def ensure_default_ct(ct, ext):
        types = {'png': 'image/png', 'jpeg': 'image/jpeg', 'jpg': 'image/jpeg', 'gif': 'image/gif',
                 'svg': 'image/svg+xml', 'bmp': 'image/bmp', 'tif': 'image/tiff', 'tiff': 'image/tiff',
                 'emf': 'image/x-emf', 'wmf': 'image/x-wmf'}
        for d in ct.findall('{%s}Default' % NS['ct']):
            if d.get('Extension', '').lower() == ext:
                return
        d = etree.SubElement(ct, '{%s}Default' % NS['ct'])
        d.set('Extension', ext)
        d.set('ContentType', types.get(ext, 'application/octet-stream'))
        # Defaults must precede Overrides
        ct.remove(d)
        first_override = ct.find('{%s}Override' % NS['ct'])
        if first_override is not None:
            first_override.addprevious(d)
        else:
            ct.append(d)

    def add_landscape_footer(self, master, rels, ct):
        """footer3.xml: the template footer band widened to the landscape page width."""
        footer = master.parts['word/footer1.xml']
        frels = etree.fromstring(master.parts['word/_rels/footer1.xml.rels'])
        img_target = frels[0].get('Target')
        img = Image.open(io_bytes(master.parts['word/' + img_target])).convert('RGBA')
        wpx, hpx = img.size
        root = etree.fromstring(footer)
        ext = root.find('.//wp:extent', NS)
        cx = int(ext.get('cx'))
        new_cx = int((PAGE_H * 635) * 1.04)
        new_wpx = int(round(wpx * new_cx / cx))
        canvas = Image.new('RGBA', (new_wpx, hpx), img.getpixel((wpx - 1, hpx // 2)))
        canvas.paste(img, (0, 0))
        buf = bytes_io()
        canvas.save(buf, 'PNG', optimize=True)
        master.set_bytes('word/media/footer_band_landscape.png', buf.getvalue())
        ext.set('cx', str(new_cx))
        for x in root.iter(w('a:ext')):
            if x.getparent().tag == w('a:xfrm'):
                x.set('cx', str(new_cx))
        for dp in root.iter(w('wp:docPr')):
            dp.set('id', '4999')
        master.set_xml('word/footer3.xml', root)
        fr = etree.Element('{%s}Relationships' % NS['rel'], nsmap={None: NS['rel']})
        r = etree.SubElement(fr, '{%s}Relationship' % NS['rel'])
        r.set('Id', 'rId1')
        r.set('Type', REL_IMAGE)
        r.set('Target', 'media/footer_band_landscape.png')
        master.set_xml('word/_rels/footer3.xml.rels', fr)
        o = etree.SubElement(ct, '{%s}Override' % NS['ct'])
        o.set('PartName', '/word/footer3.xml')
        o.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml')
        rid = 'rIdFooterLandscape'
        nr = etree.SubElement(rels, '{%s}Relationship' % NS['rel'])
        nr.set('Id', rid)
        nr.set('Type', REL_FOOTER)
        nr.set('Target', 'footer3.xml')
        return rid

    def merge_footnotes(self, master, src, elems):
        refs = [r for el in elems for r in el.iter(w('w:footnoteReference'))]
        if not refs or 'word/footnotes.xml' not in src.parts:
            return
        sfn = src.xml('word/footnotes.xml')
        mfn = master.xml('word/footnotes.xml')
        ids = [int(f.get(w('w:id'))) for f in mfn.findall(w('w:footnote'))]
        nxt = max(ids + [0]) + 1
        smap = {f.get(w('w:id')): f for f in sfn.findall(w('w:footnote'))}
        for r in refs:
            old = r.get(w('w:id'))
            f = smap.get(old)
            if f is None:
                continue
            nf = copy.deepcopy(f)
            nf.set(w('w:id'), str(nxt))
            for p in nf.iter(w('w:p')):
                ppr = ensure_ppr(p)
                if ppr.find(w('w:rPr')) is None:
                    pass
                for run in p.iter(w('w:r')):
                    rpr = run.find(w('w:rPr'))
                    if rpr is None:
                        rpr = E('w:rPr')
                        run.insert(0, rpr)
                    if rpr.find(w('w:sz')) is None:
                        rpr.append(frag(SEGOE))
                        rpr.append(E('w:sz', {'w:val': '18'}))
            fix_schema_order(nf)
            mfn.append(nf)
            r.set(w('w:id'), str(nxt))
            nxt += 1
        master.set_xml('word/footnotes.xml', mfn)

    def collect_headings(self, body, toc_head):
        heads = []
        counters = [0, 0, 0]
        n = 0
        started = False
        for el in body:
            if el is toc_head:
                started = True
                continue
            if not started or el.tag != w('w:p'):
                continue
            lvl = heading_level(el)
            if lvl < 1 or lvl > 3:
                continue
            counters[lvl - 1] += 1
            for k in range(lvl, 3):
                counters[k] = 0
            label = '.'.join(str(c) for c in counters[:lvl]) + '.'
            n += 1
            bm = '_Toc%09d' % (300000000 + n)
            # bookmark around the heading text
            first_run = next((c for c in el if c.tag in (w('w:r'), w('w:hyperlink'))), None)
            bs = E('w:bookmarkStart', {'w:id': '0', 'w:name': bm})
            be = E('w:bookmarkEnd', {'w:id': '0'})
            if first_run is not None:
                first_run.addprevious(bs)
            else:
                el.append(bs)
            el.append(be)
            heads.append({'level': lvl, 'label': label, 'text': text_of(el).strip(), 'bookmark': bm, 'el': el})
        return heads

    @staticmethod
    def renumber_bookmarks(body):
        n = 0
        mapping = {}
        for el in body.iter(w('w:bookmarkStart'), w('w:bookmarkEnd')):
            pass
        # pair starts and ends in document order using the original ids
        open_ids = {}
        for el in body.iter():
            if el.tag == w('w:bookmarkStart'):
                n += 1
                old = el.get(w('w:id'))
                open_ids.setdefault(old, []).append(n)
                el.set(w('w:id'), str(n))
            elif el.tag == w('w:bookmarkEnd'):
                old = el.get(w('w:id'))
                stack = open_ids.get(old)
                if stack:
                    el.set(w('w:id'), str(stack.pop(0)))
                else:
                    el.getparent().remove(el)
        return mapping

    @staticmethod
    def strip_custom_xml(master, rels, ct):
        for r in list(rels):
            if r.get('Type') == REL_CUSTOMXML:
                rels.remove(r)
        for name in list(master.parts):
            if name.startswith('customXml/') or name == 'docProps/custom.xml':
                master.delete(name)
        for o in list(ct.findall('{%s}Override' % NS['ct'])):
            pn = o.get('PartName', '')
            if pn.startswith('/customXml/') or pn == '/docProps/custom.xml':
                ct.remove(o)
        root_rels = master.xml('_rels/.rels')
        for r in list(root_rels):
            if r.get('Type') == REL_CUSTOMPROPS:
                root_rels.remove(r)
        master.set_xml('_rels/.rels', root_rels)

    @staticmethod
    def set_settings(master):
        s = master.xml('word/settings.xml')
        if s.find(w('w:updateFields')) is None:
            uf = E('w:updateFields', {'w:val': 'true'})
            after = ['hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids', 'mathPr',
                     'attachedSchema', 'themeFontLang', 'clrSchemeMapping', 'doNotIncludeSubdocsInStats',
                     'doNotAutoCompressPictures', 'forceUpgrade', 'captions', 'readModeInkLockDown', 'smartTagType',
                     'schemaLibrary', 'shapeDefaults', 'doNotEmbedSmartTags', 'decimalSymbol', 'listSeparator']
            placed = False
            for ch in s:
                if etree.QName(ch).localname in after:
                    ch.addprevious(uf)
                    placed = True
                    break
            if not placed:
                s.append(uf)
        # attachedSchema entries refer to the removed SharePoint customXml
        for el in s.findall(w('w:attachedSchema')):
            s.remove(el)
        master.set_xml('word/settings.xml', s)

    def set_props(self, master):
        meta = self.meta
        core = master.xml('docProps/core.xml')

        def setc(tag, val):
            el = core.find(tag, NS)
            if el is None:
                p, t = tag.split(':')
                el = etree.SubElement(core, '{%s}%s' % (NS[p], t))
            el.text = val
        now = dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
        setc('dc:title', '%s' % meta['title'])
        setc('dc:subject', meta['id'])
        setc('dc:creator', 'iorta TechNXT')
        setc('cp:lastModifiedBy', 'iorta TechNXT')
        setc('cp:keywords', '; '.join(x for x in (meta['id'], 'TASCO Insurance', 'iorta TechNXT') if x))
        setc('dc:description', meta['subtitle'])
        setc('cp:revision', '1')
        created = core.find('dcterms:created', NS)
        if created is not None:
            created.text = now
        modified = core.find('dcterms:modified', NS)
        if modified is not None:
            modified.text = now
        master.set_xml('docProps/core.xml', core)
        app = master.xml('docProps/app.xml')
        for t in ('TotalTime', 'Pages', 'Words', 'Characters', 'Lines', 'Paragraphs', 'CharactersWithSpaces'):
            el = app.find('ep:' + t, NS)
            if el is not None:
                app.remove(el)
        comp = app.find('ep:Company', NS)
        if comp is None:
            comp = etree.SubElement(app, '{%s}Company' % NS['ep'])
        comp.text = 'iorta TechNXT'
        master.set_xml('docProps/app.xml', app)

    # ------------------------------------------------------------------
    def refresh_toc(self, path):
        mode = self.args.toc_mode
        info = {'toc_mode': mode}
        if mode == 'none':
            return info
        js = os.path.join(self.work, 'toc.json')
        cmd = [UNO_PYTHON, LO_TOOLS, path, '--toc-json', js]
        rt = None
        if mode == 'roundtrip':
            rt = os.path.join(self.work, 'roundtrip.docx')
            cmd += ['--save', rt]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=900)
        if res.returncode != 0 or not os.path.isfile(js):
            warn('LibreOffice TOC refresh failed: ' + res.stderr.strip()[-300:])
            return info
        with open(js, encoding='utf-8') as f:
            data = json.load(f)
        toc = [t for t in data['toc'] if t['page'] is not None]
        heads = self.headings
        info['toc_entries_lo'] = len(toc)
        info['headings'] = len(heads)

        def norm(s):
            s = re.sub(r'^\s*[\d.]+\s*', '', s)
            return re.sub(r'\W+', '', s).lower()
        pages = [None] * len(heads)
        if len(toc) == len(heads):
            ok = sum(1 for t, h in zip(toc, heads) if norm(t['text']).endswith(norm(h['text'])[-20:]))
            info['toc_text_matches'] = ok
            for i, t in enumerate(toc):
                pages[i] = t['page']
        else:
            warn('LibreOffice TOC has %d entries, document has %d headings; matching by text' % (len(toc), len(heads)))
            j = 0
            for i, h in enumerate(heads):
                for k in range(j, len(toc)):
                    if norm(toc[k]['text']).endswith(norm(h['text'])[-20:]):
                        pages[i] = toc[k]['page']
                        j = k + 1
                        break
        # cross-check with layout heading pages
        lay = data.get('headings') or []
        if len(lay) == len(heads):
            diff = sum(1 for p, l in zip(pages, lay) if p is not None and p != l['page'])
            info['toc_vs_layout_differences'] = diff
            for i, l in enumerate(lay):
                if pages[i] is None:
                    pages[i] = l['page']
        info['toc_missing_pages'] = sum(1 for p in pages if p is None)
        info['toc'] = [{'label': h['label'], 'text': h['text'], 'page': p} for h, p in zip(heads, pages)]
        if mode == 'roundtrip' and rt and os.path.isfile(rt):
            shutil.copy(rt, path)
            return info
        # inject the page numbers into our own (template-faithful) document
        pkg = Package(path)
        doc = pkg.xml('word/document.xml')
        page_by_bm = {h['bookmark']: p for h, p in zip(heads, pages)}
        for instr in doc.iter(w('w:instrText')):
            m = re.match(r'\s*PAGEREF\s+(\S+)', instr.text or '')
            if not m or m.group(1) not in page_by_bm:
                continue
            page = page_by_bm[m.group(1)]
            if page is None:
                continue
            # the result run follows the 'separate' fldChar
            r = instr.getparent()
            nxt = r.getnext()
            while nxt is not None and nxt.find(w('w:fldChar')) is None:
                nxt = nxt.getnext()
            res_run = nxt.getnext() if nxt is not None else None
            if res_run is not None and res_run.find(w('w:t')) is not None:
                res_run.find(w('w:t')).text = str(page)
        pkg.set_xml('word/document.xml', doc)
        pkg.save(path)
        return info

    def validate(self, path):
        rep = {'output': path}
        pkg = Package(path)
        problems = []
        for name in pkg.parts:
            if name.endswith('.rels'):
                ids = [r.get('Id') for r in etree.fromstring(pkg.parts[name])]
                dup = {i for i in ids if ids.count(i) > 1}
                if dup:
                    problems.append('duplicate relationship ids in %s: %s' % (name, sorted(dup)))
        rels = {r.get('Id') for r in pkg.xml('word/_rels/document.xml.rels')}
        doc = pkg.xml('word/document.xml')
        used = {v for el in doc.iter() for k, v in el.attrib.items() if k.startswith('{%s}' % R_NS)}
        missing = used - rels
        if missing:
            problems.append('undefined relationship ids in document.xml: %s' % sorted(missing))
        dp = [d.get('id') for d in doc.iter(w('wp:docPr'))]
        if len(dp) != len(set(dp)):
            problems.append('duplicate drawing ids')
        bms = [b.get(w('w:name')) for b in doc.iter(w('w:bookmarkStart'))]
        if len(bms) != len(set(bms)):
            problems.append('duplicate bookmark names')
        ct = pkg.xml('[Content_Types].xml')
        overrides = {o.get('PartName') for o in ct.findall('{%s}Override' % NS['ct'])}
        for o in overrides:
            if o.lstrip('/') not in pkg.parts:
                problems.append('content type override for missing part %s' % o)
        try:
            import docx  # python-docx
            d = docx.Document(path)
            rep['python_docx'] = 'ok (%d paragraphs, %d tables, %d sections)' % (
                len(d.paragraphs), len(d.tables), len(d.sections))
        except Exception as e:  # pragma: no cover
            problems.append('python-docx failed to open: %s' % e)
        rep['problems'] = problems
        rep['sections'] = len(doc.findall('.//w:sectPr', NS))
        rep['images'] = len([n for n in pkg.parts if n.startswith('word/media/body_')])
        for p in problems:
            warn(p)
        return rep

    def preview(self, path, outdir):
        os.makedirs(outdir, exist_ok=True)
        base = os.path.splitext(os.path.basename(path))[0]
        pdf = os.path.join(outdir, base + '.pdf')
        res = subprocess.run([UNO_PYTHON, LO_TOOLS, path, '--pdf', pdf], capture_output=True, text=True, timeout=900)
        if res.returncode != 0 or not os.path.isfile(pdf):
            warn('preview failed: ' + res.stderr[-300:])
            return
        prefix = os.path.join(outdir, re.sub(r'[^A-Za-z0-9_-]+', '_', base))
        subprocess.run(['pdftoppm', '-r', '60', '-png', pdf, prefix], check=False)


def io_bytes(b):
    import io
    return io.BytesIO(b)


def bytes_io():
    import io
    return io.BytesIO()


def main(argv=None):
    ap = argparse.ArgumentParser(description='Markdown with YAML front matter -> iorta TechNXT Word document')
    ap.add_argument('input')
    ap.add_argument('output')
    ap.add_argument('--img-dir', help='extra folder to look up images in')
    ap.add_argument('--template', default=DEFAULT_TEMPLATE)
    ap.add_argument('--toc-mode', choices=['inject', 'roundtrip', 'none'], default='inject',
                    help='inject (default): LibreOffice computes the TOC page numbers, which are written into the '
                         'template-faithful document; roundtrip: keep LibreOffice\'s re-saved .docx; none: skip')
    ap.add_argument('--preview', metavar='DIR', help='also write a PDF and PNG page previews (LibreOffice) to DIR')
    ap.add_argument('--keep-work', action='store_true', help='keep the temporary work folder')
    ap.add_argument('--json', action='store_true', help='print the build report as JSON')
    ap.add_argument('--quiet', action='store_true')
    args = ap.parse_args(argv)
    WARNINGS.clear()
    b = Builder(args)
    rep = b.run()
    if args.json:
        print(json.dumps({k: v for k, v in rep.items() if k != 'toc'}, indent=1, ensure_ascii=False))
    elif not args.quiet:
        print('built %s  (headings %s, TOC entries from LibreOffice %s, figures %s, sections %s, problems %d, warnings %d)'
              % (args.output, rep.get('headings'), rep.get('toc_entries_lo'), rep.get('figures'),
                 rep.get('sections'), len(rep['problems']), len(rep['warnings'])))
    return 1 if rep['problems'] else 0


if __name__ == '__main__':
    sys.exit(main())
